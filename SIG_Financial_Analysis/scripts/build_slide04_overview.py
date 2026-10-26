"""Slide 4 v2: Financial Analysis - Overview, left half rebuilt for comparability (mirrors slide 3's left half).

Base = the user's own slide 4 from the OneDrive working deck (local copy made with PowerPoint SaveCopyAs).
Removes the Rp T stacked chart, its logos, the grey pointer panel and the resolved reviewer note, then adds:
  1. 100% stacked revenue split (COGS / other costs / net profit), SIG vs INTP, FY21 vs FY25
  2. "COGS share of revenue, FY21 vs. FY25" change boxes (logo + pp + arrow), as slide 3's share-change boxes
  3. FY25 compare bars (revenue, COGS, net profit), as slide 3's size comparison
The right half (NPM/ROA), title, subheadline and Key Question are left untouched.
Run:  python build_slide04_overview.py
"""
import io
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).parent))
from sig_data import check_numbers, row_values, wb  # noqa: E402
from skha_helpers import (BODY, MUTED, add_column_chart, add_legend, add_rect, add_text,  # noqa: E402
                          fill_text_frame, make_slim_deck, remove_shape, replace_runs_text, set_plot_area)

SCRATCH = Path(r"D:\Skripsi\stata_tmp\claude\d--Download-Farah-Assignment-3-Fashion-AI-fashion-ai-claude-code-starter-v4-"
               r"fashion-ai-claude-code-starter-v4\cad0ffbf-cfa8-46ee-bb48-ed5a28e089cc\scratchpad")
SRC = SCRATCH / "deck_onedrive_copy.pptx"        # SaveCopyAs of the OneDrive working deck (slide 3 = Industry)
OUT = Path(r"D:\Download\SIG_Financial_Analysis\slides\Slide_04_Overview_v2.pptx")

C_COGS, C_OTH, C_NP = "C5DAEE", "6FA8DC", "10719D"    # the user's current chart colours
C_SIG_BAR, C_INTP_BAR = "1F4E79", "9DC3E6"            # slide 3 compare bars
BOX, RED_ARROW, FLAT_ARROW, TXT_DARK, TXT_MID = "F2F2F2", "C00000", "A6A6A6", "262626", "595959"
REMOVE = ["Revenue vs cost SIG", "Picture 13", "Picture 14", "Rectangle 27", "Rectangle 7"]

# ---- data: Profitability sheet rows 32-35 (source of the current chart), Rp miliar --------
ws = wb()["Profitability"]
col = lambda row, c0: [ws.cell(row, c0 + i).value for i in range(5)]  # noqa: E731
assert ws.cell(32, 2).value == "COGS" and ws.cell(35, 2).value == "Revenue"
S = {k: col(r, 3) for k, r in (("cogs", 32), ("oth", 33), ("np", 34), ("rev", 35))}   # SIG cols C-G
I = {k: col(r, 9) for k, r in (("cogs", 32), ("oth", 33), ("np", 34), ("rev", 35))}   # INTP cols I-M
for d in (S, I):
    for k, v in d.items():
        check_numbers(k, v)
    assert all(abs(c + o + n - r) < 0.5 for c, o, n, r in zip(d["cogs"], d["oth"], d["np"], d["rev"]))
share = lambda d, key, k: d[key][k] / d["rev"][k] * 100  # noqa: E731
split = {key: [round(share(S, key, 0), 1), round(share(S, key, 4), 1), None,
               round(share(I, key, 0), 1), round(share(I, key, 4), 1)] for key in ("cogs", "oth", "np")}
d_sig = round(split["cogs"][1] - split["cogs"][0], 1)      # from the displayed shares, no visible rounding gap
d_intp = round(split["cogs"][4] - split["cogs"][3], 1)
np_mult, rev_ratio = I["np"][4] / S["np"][4], I["rev"][4] / S["rev"][4]
# COGS driver check (IS (Build) cost structure, % of revenue)
energy, overhead = row_values("IS (Build)", "Bahan bakar & energi"), row_values("IS (Build)", "Overhead lain-lain + sewa")

assert split["cogs"][:2] == [74.5, 83.3] and split["cogs"][3:] == [79.6, 80.0]
assert split["oth"][:2] == [19.9, 16.2] and split["oth"][3:] == [8.3, 11.1]
assert split["np"][:2] == [5.6, 0.5] and split["np"][3:] == [12.1, 8.9]
assert (d_sig, d_intp) == (8.8, 0.4)
assert round(np_mult) == 8 and 0.45 <= rev_ratio <= 0.55                       # "8x ... on half the revenue"
assert energy[4] - energy[0] > 0.03 and overhead[4] - overhead[0] > 0.03       # "factory overhead and energy"
fy25 = {"Revenue": (S["rev"][4], I["rev"][4]), "COGS": (S["cogs"][4], I["cogs"][4]),
        "Net profit": (S["np"][4], I["np"][4])}
DEC = {"Revenue": 1, "COGS": 1, "Net profit": 2}      # net profit is small, so two decimals
txt_t = lambda v, label: f"Rp{v / 1e3:.{DEC[label]}f} T"  # noqa: E731
assert [txt_t(v, lab) for lab, pair in fy25.items() for v in pair] == ["Rp35.2 T", "Rp17.7 T", "Rp29.3 T",
                                                                       "Rp14.2 T", "Rp0.19 T", "Rp1.58 T"]

# ---- logos from slide 3 (small wordmarks used in its share-change boxes) ---------------
full = Presentation(SRC)
s3 = full.slides[2]
assert "Industry" in " ".join(sh.text_frame.text for sh in s3.shapes if sh.has_text_frame)
logo = {}
for key, name in (("SIG", "Picture 23"), ("INTP", "Picture 25")):
    pic = next(sh for sh in s3.shapes if sh.name == name)
    # keep the source crop (the Indocement wordmark is the cropped bottom strip of a larger image)
    logo[key] = (pic.image.blob, pic.width / pic.height,
                 (pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom))
del full

# ---- slim deck with the user's slide 4 ----------------------------------------------------
prs = make_slim_deck(SRC, keep_slide_no=4)
slide = prs.slides[0]
assert "Overview" in " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
for sh in list(slide.shapes):
    if sh.name in REMOVE:
        rids = [v for el in sh._element.iter() for k, v in el.attrib.items()
                if k in (qn("r:id"), qn("r:embed"), qn("r:link"))]
        remove_shape(sh)
        for rid in rids:
            slide.part.drop_rel(rid)
group = next(sh for sh in slide.shapes if sh.name == "Group 34")
caption = next(sh for sh in group.shapes if sh.name == "Rectangle 9")
replace_runs_text(caption, "Revenue split, % of revenue, FY21 vs. FY25")


def add_logo(key, x, y, h):
    blob, ratio, (cl, cr, ct, cb) = logo[key]
    pic = slide.shapes.add_picture(io.BytesIO(blob), Pt(x), Pt(y), Pt(h * ratio), Pt(h))
    pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = cl, cr, ct, cb
    return h * ratio


# ---- row 1 left: 100% stacked revenue split ----------------------------------------------
CX, CY, CW, CH = 41.0, 194.0, 300.0, 136.0
PX0, PW, PY0, PH = 0.02, 0.96, 0.02, 0.84
cats = ["FY21", "FY25", " ", "FY21", "FY25"]
gf = add_column_chart(
    slide, CX, CY, CW, CH, cats,
    [{"name": "COGS", "values": split["cogs"], "color": C_COGS, "label_size": 9,
      "label_points": [0, 1, 3, 4]},
     {"name": "Other costs", "values": split["oth"], "color": C_OTH, "label_size": 9,
      "label_points": [0, 1, 3, 4]},
     {"name": "Net profit", "values": split["np"], "color": C_NP, "label_color": "FFFFFF", "label_size": 9,
      "label_points": [k for k, v in enumerate(split["np"]) if v is not None and v >= 8]}],
    stacked=True, gap_width=45, y_min=0, y_max=100, number_format='0.0"%"', name="Chart revenue split")
set_plot_area(gf.chart, PX0, PY0, PW, PH)
xc = [CX + CW * (PX0 + PW * (i + 0.5) / 5) for i in range(5)]
for key, (a, b) in (("SIG", (0, 1)), ("INTP", (3, 4))):
    h = 13 if key == "SIG" else 9
    w = logo[key][1] * h
    add_logo(key, (xc[a] + xc[b]) / 2 - w / 2, 178 + (13 - h) / 2, h)
add_legend(slide, CX, 334, [("COGS", C_COGS, "square"), ("Other costs", C_OTH, "square"),
                           ("Net profit", C_NP, "square")], gap=12)

# ---- row 1 right: COGS share change boxes (slide 3 share-change style) -------------------
BXX, BWW = 392.0, 128.0
add_text(slide, BXX, 180, BWW, 26, ["COGS share of revenue", "FY21 vs. FY25"], size=10, color=BODY)
for i, (key, val, arrow) in enumerate((("SIG", d_sig, RED_ARROW), ("INTP", d_intp, FLAT_ARROW))):
    by = 210 + i * 58
    add_rect(slide, BXX, by, BWW, 52, fill=BOX, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.12
    add_logo(key, BXX + 9, by + 7, 12 if key == "SIG" else 8)
    t = add_text(slide, BXX + 9, by + 22, 80, 24, [[(f"{val:+.1f}", {"bold": True, "size": 16}),
                                                   (" pp", {"size": 11})]], size=16, color=BODY, anchor="m")
    t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_rect(slide, BXX + BWW - 24, by + 29, 13, 11, fill=arrow, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
add_text(slide, BXX, 328, BWW + 4, 12, "Driven by factory overhead and energy", size=9, color=MUTED)

# ---- row 2: FY25 compare bars (slide 3 size-comparison style) ----------------------------
add_text(slide, CX, 352, 479, 14, [[("Indocement earns 8x SIG's net profit on half the revenue", {"bold": True}),
                                    (" (FY25):", {})]], size=10.5, color=BODY)
BAR_MAX = 52.0
for j, (label, (sv, iv)) in enumerate(fy25.items()):
    bx = CX + j * 160
    add_text(slide, bx, 372, 150, 14, [[(f"{label} (FY25)", {"bold": True})]], size=10, color=TXT_DARK)
    top = max(sv, iv)
    for k, (who, v, c) in enumerate((("SIG", sv, C_SIG_BAR), ("INTP", iv, C_INTP_BAR))):
        ry = 392 + k * 20
        add_text(slide, bx, ry, 28, 14, who, size=10, color=TXT_MID)
        L = max(BAR_MAX * v / top, 2.0)
        add_rect(slide, bx + 30, ry + 2, L, 10, fill=c)
        add_text(slide, bx + 30 + L + 4, ry, 70, 14, [[(txt_t(v, label), {"bold": True})]], size=10, color=TXT_DARK)
add_text(slide, CX, 433, 300, 12, "COGS includes freight for both companies", size=9, color="7F7F7F")

# ---- speaker notes (content of the removed pointer panel, now shown visually) -------------
oth_avg_s = sum(o / r for o, r in zip(S["oth"], S["rev"])) / 5 * 100
oth_avg_i = sum(o / r for o, r in zip(I["oth"], I["rev"])) / 5 * 100
notes = "\n".join([
    "Revenue split = COGS (including freight, for both companies) + other costs + net profit = 100% of revenue. "
    "Net profit is profit attributable to owners for SIG and excludes the FY25 one-off gain for INTP. Source: "
    "Profitability sheet of the working paper.",
    f"COGS rose from {split['cogs'][0]}% of revenue in FY21 to {split['cogs'][1]}% in FY25 (+{d_sig} pp), while INTP "
    f"stayed at {split['cogs'][3]}% to {split['cogs'][4]}% (+{d_intp} pp). The rise was driven by factory overhead "
    f"(+{(overhead[4] - overhead[0]) * 100:.1f} pp of revenue) and energy (+{(energy[4] - energy[0]) * 100:.1f} pp).",
    f"Other costs fell from {split['oth'][0]}% to {split['oth'][1]}% of revenue but averaged {oth_avg_s:.1f}% vs "
    f"{oth_avg_i:.1f}% at INTP.",
    f"Net profit to owners fell from Rp{S['np'][0] / 1e3:.2f} T to Rp{S['np'][4] / 1e3:.2f} T. Since FY24 INTP's net "
    f"profit has exceeded SIG's, and in FY25 it was {np_mult:.1f}x SIG's on {rev_ratio * 100:.0f}% of the revenue.",
])
slide.notes_slide.notes_text_frame.text = notes

slide_text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
assert ";" not in slide_text and ";" not in notes, "semicolon found"

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print("saved", OUT, f"{OUT.stat().st_size / 1e6:.2f} MB")
print("split", split, "d", d_sig, d_intp, "np mult", round(np_mult, 2), "rev ratio", round(rev_ratio, 3))
