"""Slide 6 v3: Leverage (solvability vs repayability) and Liquidity; Efficiency removed.

Delivers ONLY this slide as a slim .pptx (template from v1, one layout, Univers fonts only).
Run:  python build_slide6_v3.py
"""
import sys
from pathlib import Path

from pptx.enum.chart import XL_LABEL_POSITION as POS

sys.path.insert(0, str(Path(__file__).parent))
from sig_data import YEARS, check_numbers, ratio, row_values  # noqa: E402
from skha_helpers import (BODY, NAVY, RED, SKY, add_banner, add_column_chart, add_column_header,  # noqa: E402
                          add_kpi_table, add_legend, add_line, add_line_chart, add_takeaway_box, add_text,
                          clear_slide, fill_text_frame, make_slim_deck, replace_runs_text)

SRC = Path(r"D:\Download\Financial Analysis SIG (SMGR) v1.pptx")
OUT = Path(r"D:\Download\SIG_Financial_Analysis\slides\Slide_06_Leverage_Liquidity_v3.pptx")

# ---- data (all from the WP workbook) -----------------------------------------
de, de_i = ratio(27)
dta, dta_i = ratio(28)
ibd_eb, ibd_eb_i = ratio(31)
tie, tie_i = ratio(33)
eb_int, _ = ratio(34)
cur, cur_i = ratio(36)
quick, quick_i = ratio(37)
ocf_cl, ocf_cl_i = ratio(40)
cash_due, cash_due_i = ratio(41)
ibd = [v / 1e6 for v in row_values("BS (Presentation)", "Memo: Utang berbunga", startswith=False)]
ibd_due = [v / 1e6 for v in row_values("BS (Presentation)", "Memo: Utang berbunga jatuh tempo", startswith=False)]
cash = [v / 1e6 for v in row_values("BS (Presentation)", "Kas dan setara kas")]
ret_earn = [v / 1e6 for v in row_values("BS (Presentation)", "Saldo laba – belum dicadangkan")]
ebitda = [v / 1e6 for v in row_values("IS (Build)", "EBITDA (EBIT + D&A)")]
ni = [v / 1e6 for v in row_values("IS (Build)", "Laba ke pemilik entitas induk")]
fin_cost = [abs(v) / 1e6 for v in row_values("IS (Build)", "Beban keuangan")]
cfo = [v / 1e6 for v in row_values("CF (Presentation)", "Arus kas bersih dari aktivitas operasi")]
div = [abs(v) / 1e6 for v in row_values("CF (Presentation)", "Pembayaran dividen – pemilik entitas induk")]
intp_fin = abs(row_values("Peer_INTP", "Biaya keuangan")[4]) / 1e6
intp_ibd = row_values("Peer_INTP", "Utang berbunga (bank + sewa)")[4] / 1e6
intp_ni = row_values("INTP_3FS", "Laba tahun berjalan tanpa one-off")[4] / 1e6
for name, vals in dict(de=de, dta=dta, ibd_eb=ibd_eb, tie=tie, eb_int=eb_int, cur=cur, quick=quick, ocf_cl=ocf_cl,
                       cash_due=cash_due, ibd=ibd, ibd_due=ibd_due, cash=cash, ret_earn=ret_earn, ebitda=ebitda,
                       ni=ni, fin_cost=fin_cost, cfo=cfo, div=div, intp=[intp_fin, intp_ibd, intp_ni],
                       peers=[de_i[4], dta_i[4], ibd_eb_i[4], tie_i[4], cur_i[4], quick_i[4], ocf_cl_i[4],
                              cash_due_i[4]]).items():
    check_numbers(name, vals)

debt_ni = [d / n for d, n in zip(ibd, ni)]
debt_ni_i = intp_ibd / intp_ni
cfo_debt = [c / d for c, d in zip(cfo, ibd)]
chg = lambda a, b: b / a - 1  # noqa: E731
ibd_chg, ebitda_chg, fin_chg = chg(ibd[0], ibd[4]), chg(ebitda[0], ebitda[4]), chg(fin_cost[0], fin_cost[4])
# story guards: the headline claims must hold in the data
assert abs(ibd_chg - ebitda_chg) < 0.05, (ibd_chg, ebitda_chg)       # "EBITDA fell as fast as debt"
assert ibd_chg < -0.45 and de[4] < 0.3                                 # "debt halved", "only 0.24x equity"
assert abs(ibd_eb[4] - ibd_eb[0]) < 0.1                                # "debt/EBITDA is still 2.3x"


def x2(v):
    return f"{v:.2f}x"


def x1(v):
    return f"{v:.1f}x"


def pc(v):
    return f"{v * 100:.1f}%"


def idn(v, d=1):
    return f"{v:.{d}f}".replace(".", ",")


DOWN = "\u25bc"

# ---- slim deck with only slide 6 ---------------------------------------------
prs = make_slim_deck(SRC, keep_slide_no=6)
slide = prs.slides[0]
clear_slide(slide)

title = [s for s in slide.placeholders if s.placeholder_format.idx == 0][0]
replace_runs_text(title, "Leverage and liquidity are not the root problem: deleveraging kept SIG solvent, but because "
                         "earnings fell as fast as debt, repayment capacity did not improve")

COLS = [40.6, 338.0, 635.3]
CW = 283.5
KPI_W = [128.5, 51.7, 51.7, 51.6]
KPI_HDR = ["SIG FY21", "SIG FY25", "INTP FY25"]
BANNER_Y, HEAD_Y, LEGEND_Y, CHART_Y, CHART_H, KPI_Y = 131, 157, 209, 220, 102, 326

# group labels: Leverage (cols 1-2) | Liquidity (col 3)
for x0, x1_, label in [(COLS[0], COLS[1] + CW, "Leverage"), (COLS[2], COLS[2] + CW, "Liquidity")]:
    add_text(slide, x0, 112, 200, 13, [[(label, {"bold": True})]], size=10.5, color=NAVY, anchor="b")
    add_line(slide, x0, 127.5, x1_, 127.5, NAVY, 0.75)

columns = [
    {
        "banner": ("Solvability", "debt vs. what SIG owns"),
        "headline": f"Debt halved to Rp{ibd[4]:.1f} T and finance cost fell {abs(fin_chg) * 100:.0f}%; "
                    f"debt is only {de[4]:.2f}x equity",
        "caption": "Interest-bearing debt, Rp tn",
        "rows": [
            ("Debt-to-equity", [x2(de[0]), x2(de[4]), x2(de_i[4])]),
            ("Debt-to-assets", [pc(dta[0]), pc(dta[4]), pc(dta_i[4])]),
            ("Finance cost, Rp tn", [f"{fin_cost[0]:.2f}", f"{fin_cost[4]:.2f}", f"{intp_fin:.2f}"]),
        ],
    },
    {
        "banner": ("Repayability", "debt vs. what SIG earns"),
        "headline": f"EBITDA halved too, so debt/EBITDA is still {ibd_eb[4]:.1f}x and EBIT covers interest "
                    f"only {tie[4]:.1f}x",
        "caption": "EBITDA and net profit, Rp tn",
        "legend": [("EBITDA", NAVY, "square"), ("Net profit", SKY, "square")],
        "rows": [
            ("Debt \u00f7 EBITDA", [x2(ibd_eb[0]), x2(ibd_eb[4]), x2(ibd_eb_i[4])]),
            ("EBIT \u00f7 interest", [x1(tie[0]), (x1(tie[4]), DOWN), x1(tie_i[4])]),
            ("Debt \u00f7 net profit", [x1(debt_ni[0]), (x1(debt_ni[4]), DOWN), x1(debt_ni_i)]),
        ],
    },
    {
        "banner": ("Short-term coverage", "vs. what falls due"),
        "headline": f"Current assets cover short-term liabilities {cur[4]:.1f}x, a thinner buffer than "
                    f"INTP\u2019s {cur_i[4]:.1f}x",
        "caption": "Current ratio, x",
        "legend": [("SIG", NAVY, "line"), ("INTP", RED, "line")],
        "rows": [
            ("Cash \u00f7 debt due \u22641 yr", [x2(cash_due[0]), x2(cash_due[4]), x2(cash_due_i[4])]),
            ("CFO \u00f7 current liabilities", [x2(ocf_cl[0]), (x2(ocf_cl[4]), DOWN), x2(ocf_cl_i[4])]),
            ("Quick ratio", [x2(quick[0]), x2(quick[4]), x2(quick_i[4])]),
        ],
    },
]
for x, col in zip(COLS, columns):
    add_banner(slide, x, BANNER_Y, CW, 22, *col["banner"])
    add_column_header(slide, x, HEAD_Y, CW, col["headline"], col["caption"])
    if col.get("legend"):
        add_legend(slide, x, LEGEND_Y, col["legend"])
    add_kpi_table(slide, x, KPI_Y, KPI_W, KPI_HDR, col["rows"], row_h=21)

# rating note sits in the solvability column's (otherwise empty) legend slot
add_text(slide, COLS[0], LEGEND_Y, CW, 11, [[("PEFINDO rating: ", {"bold": True}),
                                             "idAA+ (Aug 2022) \u2192 idAAA (Aug 2024)"]], size=9, anchor="m")

add_column_chart(slide, COLS[0], CHART_Y, CW, CHART_H, YEARS,
                 [{"name": "Interest-bearing debt", "values": ibd, "color": NAVY}],
                 gap_width=90, y_min=0, y_max=max(ibd) * 1.2, name="Chart debt")
add_column_chart(slide, COLS[1], CHART_Y, CW, CHART_H, YEARS,
                 [{"name": "EBITDA", "values": ebitda, "color": NAVY},
                  {"name": "Net profit", "values": ni, "color": SKY}],
                 gap_width=70, y_min=0, y_max=max(ebitda) * 1.22, name="Chart earnings")
add_line_chart(slide, COLS[2], CHART_Y, CW, CHART_H, YEARS,
               [{"name": "SIG", "values": cur, "color": NAVY, "labels": {0: POS.BELOW, 4: POS.BELOW}},
                {"name": "INTP", "values": cur_i, "color": RED, "labels": {0: POS.ABOVE, 4: POS.ABOVE}}],
               number_format="0.00", y_min=0.3, y_max=3.0, name="Chart liquidity")

add_takeaway_box(slide, 40.6, 416, 878.2, 72, [
    [("Second-order effects, not the root cause. ", {"bold": True}),
     "Deleveraging kept SIG solvent (rated idAAA) and liquidity is adequate, but repayment capacity stalled because "
     "EBITDA fell as fast as debt. Restoring profitability is what restores repayability, so the deep dive focuses "
     "there."],
])

source = [s for s in slide.placeholders if s.placeholder_format.idx == 10][0]
fill_text_frame(source.text_frame, [
    "Note: interest-bearing debt = bank loans, bonds, leases and temporary syirkah funds; equity and net profit = "
    "attributable to owners. Source: SIG and INTP audited financial statements FY2021\u20132025; SIG 9M25 earnings "
    "call (PEFINDO rating); Skha analysis"], size=8, color="7F7F7F")

notes = slide.notes_slide.notes_text_frame
notes.text = "\n".join([
    f"Repayability: utang berbunga turun {abs(ibd_chg) * 100:.0f}% (Rp{idn(ibd[0])} T -> Rp{idn(ibd[4])} T), EBITDA "
    f"turun {abs(ebitda_chg) * 100:.0f}% -> debt/EBITDA tetap {idn(ibd_eb[0], 2)}x -> {idn(ibd_eb[4], 2)}x.",
    f"Basis kas juga tidak membaik: CFO / utang {cfo_debt[0] * 100:.0f}% -> {cfo_debt[4] * 100:.0f}%; EBITDA / bunga "
    f"{idn(eb_int[0])}x -> {idn(eb_int[4])}x. Yang memburuk basis laba: EBIT / bunga {idn(tie[0])}x -> "
    f"{idn(tie[4])}x; utang / laba bersih {idn(debt_ni[0])}x -> {idn(debt_ni[4])}x.",
    f"Saldo laba hanya Rp{idn(ret_earn[0])} T -> Rp{idn(ret_earn[4])} T dan turun di FY25: dividen "
    f"Rp{idn(div[4], 2)} T > laba Rp{idn(ni[4], 2)} T.",
    f"Watch item: utang berbunga jatuh tempo <=1 tahun naik dari Rp{idn(ibd_due[3])} T (FY24) ke "
    f"Rp{idn(ibd_due[4])} T (FY25) vs kas Rp{idn(cash[4])} T; SIG (9M25) menyebut current ratio turun karena "
    "tambahan pinjaman jangka pendek.",
    "Rating: PEFINDO idAA+ (Agu 2022, Agu 2023) -> idAAA (Agu 2024), reaffirm idAAA/Stable Agu 2025 "
    "(SIG earnings call 9M25, hlm. 11).",
])

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print("saved", OUT, f"{OUT.stat().st_size / 1e6:.2f} MB")
print("ibd", [round(v, 2) for v in ibd], "ebitda", [round(v, 2) for v in ebitda], "ni", [round(v, 2) for v in ni])
print("debt/ni", [round(v, 1) for v in debt_ni], "intp", round(debt_ni_i, 2), "| cfo/debt", [round(v, 3) for v in cfo_debt])
print("chg ibd/ebitda/fin", round(ibd_chg, 3), round(ebitda_chg, 3), round(fin_chg, 3))
