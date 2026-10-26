"""Slide 9: Revenue = Price x Volume (revenue mix, price vs INTP and sector PPI, volume vs INTP and market, KTW).

Delivers ONLY this slide as a slim .pptx (template from v1, one layout, Univers fonts only).
Bottom charts: index FY21 = 100 with the benchmark (sector PPI, national demand) drawn as a dashed line, in the
user's reference colours (SIG #083358, INTP #358FD1, 1.5 pt lines), legend right of the plot and
"Change 21-25" summary boxes on the far right. No per-point labels: the three index lines sit within a few
points of each other, so labels would stack and hide the PPI line.
Run:  python build_slide9_revenue.py
"""
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).parent))
from sig_data import PPI_KBLI23_Q4_YOY, YEARS, block_row, check_numbers, cumulative, row_values, wb  # noqa: E402
from skha_helpers import (BODY, GREY, INK, INTP_BLUE, MUTED, PANEL_GREY, PANEL_LIGHT, SIG_NAVY, SKY, TINT,  # noqa: E402
                          add_column_chart, add_column_header, add_legend, add_line_chart, add_rect,
                          add_summary_boxes, add_text, clear_slide, fill_text_frame, make_slim_deck,
                          replace_runs_text)

SRC = Path(r"D:\Download\Financial Analysis SIG (SMGR) v1.pptx")
OUT = Path(r"D:\Download\SIG_Financial_Analysis\slides\Slide_09_Revenue_Price_Volume_v3.pptx")

# Sector producer prices: PPI_KBLI23_Q4_YOY in sig_data.py (BPS IHP, Barang Galian Bukan Logam, sources there).
# Robustness only (speaker notes, not on the slide):
#   BPS headline CPI, Dec-on-Dec (2025 from the BPS WebAPI series in D:\Download\Indonesia_CPI_BPS)
#   BPS wholesale price index (IHPB) for building materials, Dec-on-Dec (press coverage of the BPS releases)
CPI_DEC_YOY = {2022: 0.0551, 2023: 0.0261, 2024: 0.0157, 2025: 0.0292}
IHPB_BUILDING_DEC_YOY = {2022: 0.0687, 2023: 0.0054, 2024: 0.0072, 2025: 0.0206}


def pct0(v):
    """Signed whole percent; en dash for negatives (Univers has no U+2212 glyph guarantee)."""
    return ("+" if v >= 0 else "\u2013") + f"{abs(v) * 100:.0f}%"


def national_capacity_fy25():
    """Business Context 'Demand' block: FY25 domestic demand, national export and installed capacity (Mt)."""
    ws = wb()["Business Context"]
    hdr = next(c for row in ws.iter_rows() for c in row if c.value == "Installed Capacity")
    cols = {ws.cell(hdr.row, c).value: c for c in range(1, 16) if ws.cell(hdr.row, c).value}
    for r in range(hdr.row + 1, hdr.row + 8):
        if ws.cell(r, cols["Year"]).value == "FY25":
            return {k: ws.cell(r, c).value for k, c in cols.items()}
    raise KeyError("Business Context FY25 demand row")


# ---- data (all from the WP workbook) -----------------------------------------
P, B = "IS (Presentation)", "IS (Build)"
price = block_row(P, "B. Per ton", "Net revenue", "SIG")             # Rp ribu / ton
i_price = block_row(P, "B. Per ton", "Net revenue", "INTP")
vol = [v / 1e3 for v in block_row(P, "B. Per ton", "Volume", "SIG")]  # Mt
i_vol = [v / 1e3 for v in block_row(P, "B. Per ton", "Volume", "INTP")]
cost_t = [-v for v in block_row(P, "B. Per ton", "Total delivered cost", "SIG")]
i_cost_t = [-v for v in block_row(P, "B. Per ton", "Total delivered cost", "INTP")]
rev_p = block_row(P, "A. Comparable P&L", "Net revenue", "SIG")
products = [
    ("Cement", row_values(B, "Semen"), SIG_NAVY),
    ("Clinker", row_values(B, "Terak (klinker)"), SKY),
    ("Ready-mix & precast", row_values(B, "Beton jadi dan siap pakai"), "A5DDF6"),
    ("Non-cement materials", row_values(B, "Bahan bangunan non-semen"), TINT),
    ("Others", row_values(B, "Lainnya"), "D9D9D9"),
]
revenue = row_values(B, "Pendapatan")
dom = [v / 1e3 for v in row_values(B, "Volume penjualan domestik SIG")]
reg = [v / 1e3 for v in row_values(B, "Volume penjualan regional/ekspor SIG")]
mkt = [v / 1e3 for v in row_values(B, "Permintaan semen domestik nasional")]
share = row_values(B, "Pangsa pasar domestik SIG")
fixed = [sum(-row_values(B, lab)[i] for lab in ("Tenaga kerja", "Penyusutan, amortisasi & penurunan nilai", "Sewa",
                                                "Lain-lain (overhead pabrik)")) for i in range(5)]
cap = national_capacity_fy25()
for name, vals in [("price", price), ("i_price", i_price), ("vol", vol), ("i_vol", i_vol), ("cost_t", cost_t),
                   ("i_cost_t", i_cost_t), ("rev_p", rev_p), ("revenue", revenue), ("dom", dom), ("reg", reg),
                   ("mkt", mkt), ("share", share), ("fixed", fixed),
                   ("cap", [cap["Domestic Demand"], cap["National Export"], cap["Installed Capacity"]])] + \
                  [(p[0], p[1]) for p in products]:
    check_numbers(name, vals)

ppi = cumulative(PPI_KBLI23_Q4_YOY)
ppi_path = [cumulative(PPI_KBLI23_Q4_YOY, tuple(range(2022, 2022 + k))) for k in range(5)]   # FY21 = 1.0
ppi_since22 = cumulative(PPI_KBLI23_Q4_YOY, (2023, 2024, 2025))
price_chg, i_price_chg = price[4] / price[0] - 1, i_price[4] / i_price[0] - 1
rel, i_rel = (1 + price_chg) / ppi - 1, (1 + i_price_chg) / ppi - 1
rel_cpi = (1 + price_chg) / cumulative(CPI_DEC_YOY) - 1
rel_ihpb = (1 + price_chg) / cumulative(IHPB_BUILDING_DEC_YOY) - 1
peak_drop, i_peak_drop = price[4] / price[1] - 1, i_price[4] / i_price[1] - 1
vol_chg, i_vol_chg, mkt_chg = vol[4] / vol[0] - 1, i_vol[4] / i_vol[0] - 1, mkt[4] / mkt[0] - 1
share_eff = mkt[4] * (share[4] - share[0])
mkt_eff = share[0] * (mkt[4] - mkt[0])
share_part = share_eff / (dom[4] - dom[0])
fpt = [f / v / 1e3 for f, v in zip(fixed, vol)]          # Rp juta / ribu ton -> Rp ribu per ton
fpt_chg = fpt[4] / fpt[0] - 1
clinker = products[1][1]
clinker_share = [c / r for c, r in zip(clinker, revenue)]
clinker_yoy = clinker[4] / clinker[3] - 1
reg_offset = (reg[4] - reg[3]) / -(dom[4] - dom[3])
idle = 1 - (cap["Domestic Demand"] + cap["National Export"]) / cap["Installed Capacity"]
cost_chg, i_cost_chg = cost_t[4] / cost_t[0] - 1, i_cost_t[4] / i_cost_t[0] - 1

# guards: every claim on the slide must hold in the data
assert all(abs(p * v * 1e3 - r) / r < 0.001 for p, v, r in zip(price, vol, rev_p))   # per-ton block is consistent
assert all(abs(a - b) / b < 0.001 for a, b in zip(rev_p, revenue))                    # same revenue as IS (Build)
assert 0.06 < price_chg < 0.08 and 0.07 < i_price_chg < 0.09        # "+7%", "INTP +8%"
assert 0.115 < ppi - 1 < 0.125                                       # "sector producer prices +12%"
assert price_chg < ppi - 1 and i_price_chg < ppi - 1                 # "both trailed"
assert round(abs(rel) * 100) == 4 and round(abs(i_rel) * 100) == 4   # "fell 4% in real terms", "both fell about 4%"
assert max(rel, rel_cpi, rel_ihpb) < 0                               # same direction under every deflator
assert -0.12 < vol_chg < -0.10 and 0.10 < i_vol_chg < 0.12           # "fell 11%", "INTP grew 11%"
assert round(abs(mkt_chg) * 100) == 3 and mkt_chg < 0                # "the market slipped only 3%"
assert 0.65 < share_part < 0.75                                       # "71% ... lost market share"
assert 0.25 < fpt_chg < 0.35                                          # "fixed cost per ton is up 30%"
assert 0.33 < idle < 0.37                                             # "~35% of national capacity idle"
assert clinker_share[4] == max(clinker_share)                         # "five-year high"
assert 0.65 < reg_offset < 0.85                                       # "~75% of the domestic drop"

# ---- slim deck with only slide 9 ---------------------------------------------
prs = make_slim_deck(SRC, keep_slide_no=9)
slide = prs.slides[0]
clear_slide(slide)
title = [s for s in slide.placeholders if s.placeholder_format.idx == 0][0]
replace_runs_text(title, "Revenue = Price x Volume")

LX, LW = 40.6, 690.9
PX, VX, BW = 40.6, 393.5, 338.0
KX, KW = 742.0, 177.5

# ---- top: revenue by product -------------------------------------------------
add_column_header(slide, LX, 114, LW,
                  f"Revenue fell {abs(revenue[4] / revenue[0] - 1) * 100:.0f}% to Rp{revenue[4] / 1e6:.1f} T as cement "
                  f"sales shrank, and clinker's share rose to a five-year high of {clinker_share[4] * 100:.1f}%",
                  "Revenue by product, Rp tn", headline_h=16)
add_legend(slide, LX, 152, [(nm, col, "square") for nm, _, col in products], gap=9)
add_column_chart(
    slide, LX, 164, 430, 135, YEARS,
    [{"name": nm, "values": [v / 1e6 for v in vals], "color": col,
      "label_all": nm in ("Cement", "Clinker"), "label_color": "FFFFFF" if nm == "Cement" else BODY,
      "label_size": 9 if nm == "Cement" else 8}
     for nm, vals, col in products],
    stacked=True, gap_width=55, y_min=0, y_max=44,
    totals=[f"{r / 1e6:.1f}" for r in revenue], total_pad=3.2, name="Chart revenue mix")

# callout: limestone -> clinker -> cement
CX, CY, CW, CH = 482.0, 152.0, 249.5, 147.0
add_rect(slide, CX, CY, CW, CH, fill=PANEL_LIGHT)
add_text(slide, CX + 8, CY + 5, CW - 16, 13, [[("Interesting: a shift to semi-finished clinker", {"bold": True})]],
         size=10.5, color=INK)
chev_w, chev_y = 80.0, CY + 22
for i, (label, fill, color) in enumerate([("Limestone", TINT, BODY), ("Clinker", SKY, INK),
                                          ("Cement", SIG_NAVY, "FFFFFF")]):
    shp = add_rect(slide, CX + 8 + i * (chev_w - 4), chev_y, chev_w, 20, fill=fill,
                   shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
    fill_text_frame(shp.text_frame, [[(label, {"bold": True})]], size=9, color=color, align="c")
add_text(slide, CX + 8, CY + 48, CW - 16, CH - 52, [
    {"runs": [f"Clinker revenue {pct0(clinker_yoy)} in FY25 to Rp{clinker[4] / 1e6:.1f} T, "
              f"{clinker_share[4] * 100:.1f}% of revenue (highest in five years)"], "bullet": True, "space_after": 2},
    {"runs": [f"Regional & export volume +{reg[4] - reg[3]:.1f} Mt in FY25 offset ~{reg_offset * 100:.0f}% of the "
              f"{abs(dom[4] - dom[3]):.1f} Mt domestic drop, keeping kilns running"], "bullet": True, "space_after": 2},
    {"runs": [("Trade-off: ", {"bold": True}),
              "clinker is sold before grinding and bagging, so each ton carries less value added than cement"],
     "bullet": True},
], size=9)

# ---- bottom: price and volume, index FY21 = 100 with the benchmark as a dashed line ----
CH_Y, CH_H, CH_W, BOX_W = 360.0, 128.0, 204.0, 62.0


def bottom_panel(x, headline, caption, sig_vals, intp_vals, bench_vals, bench_name, y_min, y_max, name):
    add_column_header(slide, x, 309, BW, headline, caption)
    idx = lambda s: [v / s[0] * 100 for v in s]  # noqa: E731
    # benchmark first so SIG and INTP are drawn on top of it
    add_line_chart(slide, x, CH_Y, CH_W, CH_H, YEARS,
                   [{"name": bench_name, "values": idx(bench_vals), "color": MUTED, "width": 1.5, "dash": True,
                     "marker": False},
                    {"name": "INTP", "values": idx(intp_vals), "color": INTP_BLUE, "width": 1.5},
                    {"name": "SIG", "values": idx(sig_vals), "color": SIG_NAVY, "width": 1.5}],
                   number_format="0", y_min=y_min, y_max=y_max, name=name, baseline=100)
    add_legend(slide, x + CH_W + 4, CH_Y + 34,
               [("SIG", SIG_NAVY, "line"), ("INTP", INTP_BLUE, "line"), (bench_name, MUTED, "dash")],
               line_w=1.5, vertical=True)
    chg = lambda s: s[4] / s[0] - 1  # noqa: E731
    add_summary_boxes(slide, x + BW - BOX_W, CH_Y + 4, BOX_W, "Change 21\u201325", [
        ("SIG", pct0(chg(sig_vals)), SIG_NAVY, "FFFFFF", None),
        ("INTP", pct0(chg(intp_vals)), INTP_BLUE, "FFFFFF", None),
        (bench_name, pct0(chg(bench_vals)), PANEL_GREY, BODY, (MUTED, True)),
    ])


bottom_panel(PX,
             f"Price per ton rose {price_chg * 100:.0f}%, in line with INTP ({pct0(i_price_chg)}), but fell "
             f"{abs(rel) * 100:.0f}% in real terms as sector producer prices rose {(ppi - 1) * 100:.0f}%",
             "Net revenue per ton, index FY21 = 100", price, i_price, ppi_path, "Sector PPI", 97, 117,
             "Chart price index")
bottom_panel(VX,
             f"Volume fell {abs(vol_chg) * 100:.0f}% while INTP grew {i_vol_chg * 100:.0f}% and the market slipped "
             f"only {abs(mkt_chg) * 100:.0f}%, so SIG mainly lost share",
             "Sales volume, index FY21 = 100 (INTP incl. Bosowa, Grobogan)", vol, i_vol, mkt, "Market", 84, 117,
             "Chart volume index")

# ---- right: KTW --------------------------------------------------------------
add_rect(slide, KX, 114, KW, 380, fill=PANEL_GREY, line=GREY, line_w=0.75, dash=True)
add_text(slide, KX + 9, 122, KW - 18, 364, [
    {"runs": [("Key takeaways", {"bold": True, "size": 12, "color": INK})], "space_after": 10},
    {"runs": [("1. Price: nothing to fix. ", {"bold": True}),
              f"SIG kept pace with INTP ({pct0(price_chg)} vs {pct0(i_price_chg)}). Both fell about "
              f"{abs(rel) * 100:.0f}% in real terms because sector producer prices rose {(ppi - 1) * 100:.0f}%. With "
              f"~{idle * 100:.0f}% of national capacity idle, the market sets the price for every producer."],
     "space_after": 10},
    {"runs": [("2. Volume is the lever. ", {"bold": True}),
              f"Lost market share explains {share_part * 100:.0f}% of SIG's domestic volume drop, and fixed cost per "
              f"ton is up {fpt_chg * 100:.0f}% since FY21. Regain domestic share or grow exports to refill idle "
              "capacity. Clinker exports lift utilization, but each ton carries less value added than cement."]},
], size=10.5, anchor="m")

source = [s for s in slide.placeholders if s.placeholder_format.idx == 10][0]
# wider than the layout default so the note stays on two lines (skha logo starts at ~818 pt). Position and
# height are copied from the layout first, because an xfrm holding only a size would move the box to 0,0.
_l, _t, _h = source.left, source.top, source.height
source.left, source.top, source.width, source.height = _l, _t, Pt(790) - _l, _h
fill_text_frame(source.text_frame, [
    "Note: price = net revenue \u00f7 sales volume (cement & clinker). Real = deflated by sector PPI (BPS producer "
    "price index, non-metallic mineral products, KBLI 23, Q4 to Q4). Market = national cement demand (ASI). Source: "
    "SIG and INTP financial statements and annual reports FY2021\u20132025, BPS, ASI, Skha analysis"],
    size=8, color="7F7F7F")

# ---- speaker notes -------------------------------------------------------------
notes = "\n".join([
    f"Price = net revenue \u00f7 total sales volume (cement and clinker, incl. regional and export), IS (Presentation) "
    f"per-ton block. SIG Rp{price[0]:.0f}k \u2192 Rp{price[4]:.0f}k ({pct0(price_chg)}), INTP Rp{i_price[0]:.0f}k "
    f"\u2192 Rp{i_price[4]:.0f}k ({pct0(i_price_chg)}).",
    "Sector PPI = BPS producer price index for non-metallic mineral products (KBLI 23), Q4 vs Q4: "
    + ", ".join(f"{y} +{v * 100:.2f}%" for y, v in PPI_KBLI23_Q4_YOY.items())
    + f" (2022 on the 2010 base, chain-linked to the 2016 base). Cumulative +{(ppi - 1) * 100:.1f}%. Relative to it "
      f"SIG is {rel * 100:.1f}% and INTP {i_rel * 100:.1f}%.",
    "Boundary: 'real' here means deflated by the sector PPI. That index includes SIG and INTP's own cement plus glass, "
    "ceramics and concrete products, so it is a sector price benchmark, not a consumer deflator. Revenue per ton is an "
    "annual average and includes mix (clinker, export, fighting brands).",
    f"Robustness: CPI (+{(cumulative(CPI_DEC_YOY) - 1) * 100:.1f}%) gives SIG {rel_cpi * 100:.1f}%. BPS wholesale "
    f"building-material prices (+{(cumulative(IHPB_BUILDING_DEC_YOY) - 1) * 100:.1f}%) give {rel_ihpb * 100:.1f}%. "
    "Same direction under every deflator.",
    f"Pattern: price peaked in FY22 on coal pass-through, then fell three years in a row (SIG {peak_drop * 100:.1f}%, "
    f"INTP {i_peak_drop * 100:.1f}% since FY22) while the PPI rose +{(ppi_since22 - 1) * 100:.1f}%.",
    f"Volume: SIG domestic {dom[4] - dom[0]:.1f} Mt = market effect {mkt_eff:.1f} Mt + share effect {share_eff:.1f} Mt "
    f"({share_part * 100:.0f}%). Domestic share {share[0] * 100:.1f}% \u2192 {share[4] * 100:.1f}%.",
    f"Market = ASI national demand {mkt[0]:.1f} \u2192 {mkt[4]:.1f} Mt ({mkt_chg * 100:.1f}%). The Kemenperin series "
    "(INTP public expose 2026, p5) shows 63.9 \u2192 62.4 Mt (-2.3%). Either way the market slipped about 3%.",
    "INTP volume includes Semen Bosowa volume sold by INTP and Semen Grobogan, so part of its +11% is inorganic.",
    f"Idle capacity FY25: 1 - (domestic demand {cap['Domestic Demand']:.1f} + export {cap['National Export']:.1f}) "
    f"\u00f7 capacity {cap['Installed Capacity']:.1f} Mt = {idle * 100:.0f}%.",
    f"For the COGS slide: delivered cost per ton SIG {pct0(cost_chg)} vs price {pct0(price_chg)}, INTP "
    f"{pct0(i_cost_chg)} vs {pct0(i_price_chg)}.",
])
slide.notes_slide.notes_text_frame.text = notes

# user preference: no semicolons anywhere on the slide (titles, headlines, KTW, notes)
slide_text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
assert ";" not in slide_text and ";" not in notes, "semicolon found"

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print("saved", OUT, f"{OUT.stat().st_size / 1e6:.2f} MB")
print("price", [round(v) for v in price], "intp", [round(v) for v in i_price], "chg", round(price_chg, 4),
      round(i_price_chg, 4))
print("ppi cum", round(ppi - 1, 4), "rel", round(rel, 4), round(i_rel, 4), "cpi rel", round(rel_cpi, 4),
      "ihpb rel", round(rel_ihpb, 4))
print("vol", [round(v, 1) for v in vol], "intp", [round(v, 1) for v in i_vol], "chg", round(vol_chg, 4),
      round(i_vol_chg, 4), "mkt", round(mkt_chg, 4))
print("share part", round(share_part, 3), "fixed/t chg", round(fpt_chg, 3), "idle", round(idle, 3),
      "clinker share", [round(v, 3) for v in clinker_share])
print("cost/t chg", round(cost_chg, 3), round(i_cost_chg, 3), "peak drop", round(peak_drop, 3),
      round(i_peak_drop, 3))
