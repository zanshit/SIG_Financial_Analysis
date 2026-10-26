"""Read SIG / INTP figures straight from the WP workbook, so no number on a slide is hand-typed."""
from pathlib import Path

import openpyxl

XL = Path(r"D:\Download\WP - Financial Analysis SIG (SMGR) v0.xlsx")
YEARS = ["FY21", "FY22", "FY23", "FY24", "FY25"]
_wb = None

# Sector producer prices (benchmark shared by slides 9 and 10): BPS Indeks Harga Produsen (IHP),
# Industri Barang Galian Bukan Logam (KBLI 23), y-on-y Q4 vs Q4, Tabel 4 of each BRS, chain-linked:
#   2022: BRS No. 14/02/Th. XXVI (1 Feb 2023), "Industri Barang Mineral Bukan Logam", 2010=100, 156.13 -> 163.95
#   2023, 2024: BRS IHP Triwulan IV-2024 (3 Feb 2025), 2016=100, Q4-23 115.66, Q4-24 117.87
#   2025: BRS No. 15/02/Th. XXIX (2 Feb 2026), 2016=100, 117.87 -> 120.05
PPI_KBLI23_Q4_YOY = {2022: 0.0501, 2023: 0.0293, 2024: 0.0191, 2025: 0.0185}


def cumulative(yoy, years=(2022, 2023, 2024, 2025)):
    """Compound a {year: y-o-y rate} dict over `years` (FY21 = 1.0)."""
    c = 1.0
    for y in years:
        c *= 1 + yoy[y]
    return c


def wb():
    global _wb
    if _wb is None:
        _wb = openpyxl.load_workbook(XL, data_only=True)
    return _wb


def ratio(no):
    """Ratios row by its 'No' column: returns (SIG FY21-25, INTP FY21-25).
    The sheet was called 'Ratios v2' until the user renamed it to 'Ratios' (same layout)."""
    ws = wb()["Ratios"] if "Ratios" in wb().sheetnames else wb()["Ratios v2"]
    for r in range(5, 80):
        v = ws.cell(r, 1).value
        if isinstance(v, (int, float)) and int(v) == no:
            return [ws.cell(r, c).value for c in range(5, 10)], [ws.cell(r, c).value for c in range(12, 17)]
    raise KeyError(no)


def row_values(sheet, label, n=5, start_offset=1, startswith=True):
    """First row whose label matches; returns the next n numeric cells to its right."""
    ws = wb()[sheet]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and (c.value.strip().startswith(label) if startswith else label in c.value):
                vals, col = [], c.column + start_offset
                while len(vals) < n and col < c.column + 20:
                    v = ws.cell(c.row, col).value
                    if isinstance(v, (int, float)):
                        vals.append(v)
                    col += 1
                if len(vals) == n:
                    return vals
    raise KeyError(f"{sheet}: {label}")


def block_row(sheet, block, label, side="SIG"):
    """Row `label` inside the block whose header starts with `block` (e.g. 'B. Per ton' -> 'Net revenue').
    IS (Presentation) layout: label, SIG FY21-25, CAGR, spacer, INTP FY21-25 -> SIG offset 1, INTP offset 8."""
    ws = wb()[sheet]
    head = next((c for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.strip().startswith(block)), None)
    if head is None:
        raise KeyError(f"{sheet}: block {block}")
    off = {"SIG": 1, "INTP": 8}[side]
    for r in range(head.row + 1, head.row + 40):
        v = ws.cell(r, head.column).value
        if isinstance(v, str) and v.strip().startswith(label):
            return [ws.cell(r, head.column + off + i).value for i in range(5)]
    raise KeyError(f"{sheet}: {block} / {label}")


def header_block(sheet, header):
    """The 5 values directly under a header cell."""
    ws = wb()[sheet]
    for row in ws.iter_rows():
        for c in row:
            if c.value == header:
                return [ws.cell(c.row + 1 + i, c.column).value for i in range(5)]
    raise KeyError(header)


def check_numbers(name, values):
    flat = values if isinstance(values, (list, tuple)) else [values]
    assert all(isinstance(x, (int, float)) for x in flat), (name, values)
    return values
