"""Skha template helpers for python-pptx (style spec: ../Skha_Deck_Style_Guide.md).

All coordinates are in points on the 960 x 540 pt canvas.
"""
import copy

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.chart import (XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE,
                             XL_TICK_LABEL_POSITION, XL_TICK_MARK)
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

# ---- tokens -----------------------------------------------------------------
INK = "0F1F33"
NAVY = "12274C"        # SIG
SKHA_BLUE = "1E64B0"
MID_BLUE = "368FCE"
SKY = "69C6F0"
RED = "C00000"         # INTP (comparator)
SIG_NAVY = "083358"    # SIG in the user's reference charts (slide 9 onward)
INTP_BLUE = "358FD1"   # INTP in the user's reference charts
BODY = "161616"
MUTED = "7F7F7F"
AXIS = "595959"
BASELINE = "CCDCF0"
PANEL = "E1F4FC"
PANEL_LIGHT = "F0FAFE"
PANEL_GREY = "F2F2F2"
HAIRLINE = "E5E5E5"
GREY = "BFBFBF"
IDLE = "D9D9D9"
TINT = "CBEBF8"
WATCH = "FFEFAC"       # highlight palette, yellow tint

FONT_BODY = "Univers Light"
FONT_HEAD = "Univers"

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---- text -------------------------------------------------------------------
def style_run(run, size=12, color=BODY, bold=False, italic=False, font=FONT_BODY):
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)


def _set_bullet(paragraph, indent=10):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Pt(indent)))
    pPr.set("indent", str(-Pt(indent)))
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", "•")


def _no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


def fill_text_frame(tf, paras, size=12, color=BODY, align="l", space_after=0, line_spacing=None):
    """paras: list of paragraphs; a paragraph is a str, a list of runs, or a dict
    {"runs": [...], "bullet": bool, "align": "l", "space_before": pt, "size": pt}.
    A run is a str or (text, {"bold":, "color":, "size":, "italic":, "font":})."""
    tf.clear()
    first = True
    for para in paras:
        spec = para if isinstance(para, dict) else {"runs": para}
        runs = spec["runs"]
        if isinstance(runs, str):
            runs = [runs]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[spec.get("align", align)]
        if spec.get("space_before") is not None:
            p.space_before = Pt(spec["space_before"])
        p.space_after = Pt(spec.get("space_after", space_after))
        if line_spacing:
            p.line_spacing = line_spacing
        if spec.get("bullet"):
            _set_bullet(p, spec.get("indent", 10))
        else:
            _no_bullet(p)
        for r in runs:
            text, opt = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = text
            style_run(run, size=opt.get("size", spec.get("size", size)), color=opt.get("color", color),
                      bold=opt.get("bold", False), italic=opt.get("italic", False),
                      font=opt.get("font", FONT_BODY))
    return tf


def add_text(slide, x, y, w, h, paras, size=12, color=BODY, align="l", anchor="t", inset=0,
             name=None, wrap=True, space_after=0):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Pt(inset))
    tf.vertical_anchor = ANCHOR[anchor]
    if isinstance(paras, str):
        paras = [paras]
    fill_text_frame(tf, paras, size=size, color=color, align=align, space_after=space_after)
    return tb


def set_shape_text(shape, paras, size=12, color=BODY, align="l"):
    """Replace a shape's text, keeping its bodyPr (insets/anchor)."""
    fill_text_frame(shape.text_frame, paras, size=size, color=color, align=align)


def replace_runs_text(shape, text):
    """Replace the text of the first run and drop the rest, preserving run formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


# ---- shapes -----------------------------------------------------------------
def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, name=None, shape=MSO_SHAPE.RECTANGLE,
             dash=False):
    sp = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        sp.name = name
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(line_w)
        if dash:
            sp.line.dash_style = 4  # dash
    else:
        sp.line.fill.background()
    # drop theme effects (shadow) from the default style
    style = sp._element.find(qn("p:style"))
    if style is not None:
        sp._element.remove(style)
    tf = sp.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Pt(0))
    return sp


def add_line(slide, x1, y1, x2, y2, color=SKHA_BLUE, width=1.5, name=None, head=None, tail=None):
    """Straight connector; head/tail = arrowhead type at the start/end (e.g. 'triangle'), None = plain."""
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    if name:
        ln.name = name
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    style = ln._element.find(qn("p:style"))
    if style is not None:
        ln._element.remove(style)
    a_ln = ln.line._get_or_add_ln()
    for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
        if kind:
            end = etree.SubElement(a_ln, qn(tag))
            end.set("type", kind)
            end.set("w", "med")
            end.set("len", "med")
    return ln


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


# ---- composite components -------------------------------------------------
def add_column_header(slide, x, y, w, headline, caption, headline_h=30):
    """Headline (12 bold) -> 1.5 pt Skha-blue rule -> caption (10.5). Returns y below caption."""
    add_text(slide, x, y, w, headline_h, [[(headline, {"bold": True})]], size=12, anchor="b")
    rule_y = y + headline_h + 3
    add_line(slide, x, rule_y, x + w, rule_y, SKHA_BLUE, 1.5)
    add_text(slide, x, rule_y + 2.8, w, 13, caption, size=10.5)
    return rule_y + 2.8 + 13


def add_banner(slide, x, y, w, h, label, right_label=None):
    add_rect(slide, x, y, w, h, fill=NAVY)
    add_text(slide, x + 8, y, w * 0.55, h, [[(label, {"bold": True})]], size=12, color="FFFFFF", anchor="m")
    if right_label:
        add_text(slide, x + w * 0.4, y, w * 0.6 - 8, h, right_label, size=10, color=SKY, align="r", anchor="m")


def add_takeaway_box(slide, x, y, w, h, paras, size=12, inset=10):
    box = add_rect(slide, x, y, w, h, fill=PANEL)
    tf = box.text_frame
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Pt(inset))
    tf.margin_top = tf.margin_bottom = Pt(6)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_text_frame(tf, paras, size=size)
    return box


def add_legend(slide, x, y, items, size=9, gap=14, line_w=2.25, vertical=False, row_h=14):
    """items: [(label, color, kind)] with kind 'line', 'dash' or 'square'.
    Laid out left to right, or top to bottom with vertical=True (reference-chart style, right of the plot)."""
    cx, cy = x, y
    for label, color, kind in items:
        if kind == "line":
            add_line(slide, cx, cy + 5.5, cx + 14, cy + 5.5, color, line_w)
            add_rect(slide, cx + 4.5, cy + 3, 5, 5, fill=color, shape=MSO_SHAPE.OVAL)
            tx = cx + 18
        elif kind == "dash":
            ln = add_line(slide, cx, cy + 5.5, cx + 14, cy + 5.5, color, 1.5)
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            tx = cx + 18
        else:
            add_rect(slide, cx, cy + 2.5, 6, 6, fill=color)
            tx = cx + 10
        tw = len(label) * size * 0.47 + 4
        add_text(slide, tx, cy, tw, 11, label, size=size, anchor="m", wrap=False)
        if vertical:
            cy += row_h
        else:
            cx = tx + tw + gap


def add_summary_boxes(slide, x, y, w, heading, boxes, box_h=27, gap=4, heading_h=13, heading_pad=14):
    """Reference-chart summary column: bold Skha-blue heading, then rounded boxes with a small label
    line and a bold value line. boxes: [(label, value, fill, text_color, outline)], outline = None or
    (line_color, dashed). The heading may overhang the boxes by heading_pad on each side so it stays
    on one line. Returns y below the last box."""
    add_text(slide, x - heading_pad, y, w + 2 * heading_pad, heading_h, [[(heading, {"bold": True})]], size=10,
             color=SKHA_BLUE, align="c", anchor="b", wrap=False)
    cy = y + heading_h + 3
    for label, value, fill, text_color, outline in boxes:
        line, dash = outline if outline else (None, False)
        box = add_rect(slide, x, cy, w, box_h, fill=fill, line=line, dash=dash, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        box.adjustments[0] = 0.18
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_text_frame(tf, [[(label, {"size": 8.5})], [(value, {"bold": True, "size": 11})]],
                        color=text_color, align="c")
        cy += box_h + gap
    return cy - gap


def add_text_table(slide, x, y, cols, rows, header_h=15, row_h=26, size=8.5, header_fill=NAVY, group_breaks=(),
                   chip_col=None, chips=None):
    """Text table in the template's parameter-table style (P10): navy header tiles, rows split by #E5E5E5
    hairlines, group separators #C2C2C2 0.75 pt, no vertical lines.
    cols: [(header, width)]. rows: [[cell, ...]], a cell is a str or a list of runs (text, {opts}).
    chip_col: column drawn as a rounded chip, chips = {text: (fill, text_color, line_or_None)}.
    group_breaks: row indexes that start a new group. Returns y below the table."""
    total_w = sum(w for _, w in cols)
    cx = x
    for i, (hd, w) in enumerate(cols):
        tile = add_rect(slide, cx, y, w - (2 if i < len(cols) - 1 else 0), header_h, fill=header_fill)
        tile.text_frame.margin_left = Pt(4)
        tile.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_text_frame(tile.text_frame, [[(hd, {"bold": True})]], size=size, color="FFFFFF")
        cx += w
    cy = y + header_h
    for r, row in enumerate(rows):
        if r in group_breaks:
            add_line(slide, x, cy, x + total_w, cy, "C2C2C2", 0.75)
        cx = x
        for i, ((_, w), cell) in enumerate(zip(cols, row)):
            if chip_col is not None and i == chip_col:
                fill, tcol, line = chips[cell]
                cw = min(54.0, w - 8)
                chip = add_rect(slide, cx + 4, cy + (row_h - 13) / 2, cw, 13, fill=fill, line=line,
                                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                chip.adjustments[0] = 0.5
                chip.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                fill_text_frame(chip.text_frame, [[(cell, {"bold": True})]], size=size, color=tcol, align="c")
            else:
                add_text(slide, cx + 4, cy + 1, w - 8, row_h - 2, [cell] if isinstance(cell, list) else cell,
                         size=size, anchor="m")
            cx += w
        cy += row_h
        if r < len(rows) - 1 and (r + 1) not in group_breaks:
            add_line(slide, x, cy, x + total_w, cy, HAIRLINE, 0.5)
    add_line(slide, x, cy, x + total_w, cy, "C2C2C2", 0.75)
    return cy


def add_kpi_table(slide, x, y, col_w, header, rows, row_h=26, header_h=14, label_size=10, value_size=10.5,
                  bold_col=1):
    """Shape-built mini table. col_w: [label_w, v1_w, ...]; header: labels for value columns.
    rows: [(label, [values...])]. Values are strings. Returns y below the table."""
    cx = x + col_w[0]
    for i, hd in enumerate(header):
        add_text(slide, cx, y, col_w[i + 1], header_h, hd, size=8.25, color=MUTED, align="r", anchor="b")
        cx += col_w[i + 1]
    total_w = sum(col_w)
    cy = y + header_h + 2
    add_line(slide, x, cy, x + total_w, cy, GREY, 0.75)
    for label, values in rows:
        add_text(slide, x, cy, col_w[0] - 4, row_h, label, size=label_size, anchor="m")
        cx = x + col_w[0]
        for i, v in enumerate(values):
            runs = [(v, {"bold": i == bold_col})]
            if isinstance(v, tuple):  # (text, marker) e.g. ("1.5x", "▼") for a deterioration flag
                runs = [(v[0], {"bold": i == bold_col}), (" " + v[1], {"font": "Arial", "size": 7})]
            add_text(slide, cx, cy, col_w[i + 1], row_h, [runs], size=value_size, align="r", anchor="m")
            cx += col_w[i + 1]
        cy += row_h
        add_line(slide, x, cy, x + total_w, cy, HAIRLINE, 0.5)
    return cy


# ---- slim delivery ------------------------------------------------------------
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def make_slim_deck(src, keep_slide_no, keep_fonts=("Univers", "Univers Light")):
    """Open `src` and keep only slide `keep_slide_no` (1-based), its layout, the master and the
    embedded fonts in `keep_fonts`. Parts no longer reachable are not written on save, so the
    file shrinks to the template essentials. Returns the Presentation."""
    from pptx import Presentation

    prs = Presentation(src)
    sld_id_lst = prs.slides._sldIdLst
    kept_ids = set()
    for idx, sld_id in enumerate(list(sld_id_lst), 1):
        if idx == keep_slide_no:
            kept_ids.add(sld_id.get("id"))
            continue
        rid = sld_id.rId
        sld_id_lst.remove(sld_id)
        prs.part.drop_rel(rid)
    # sections (p14:sectionLst) must not point at deleted slides
    pres_el = prs.part._element
    for sec_lst in pres_el.iter(f"{{{P14}}}sectionLst"):
        for section in list(sec_lst):
            ids = section.find(f"{{{P14}}}sldIdLst")
            if ids is not None:
                for s in list(ids):
                    if s.get("id") not in kept_ids:
                        ids.remove(s)
            if ids is None or len(ids) == 0:
                sec_lst.remove(section)
    # layouts not used by the kept slide
    for master in prs.slide_masters:
        for layout in list(master.slide_layouts):
            if not layout.used_by_slides:
                master.slide_layouts.remove(layout)
    # embedded fonts
    efl = pres_el.find(qn("p:embeddedFontLst"))
    if efl is not None:
        for ef in list(efl):
            face = ef.find(qn("p:font")).get("typeface")
            if face not in keep_fonts:
                rids = [ch.get(qn("r:id")) for ch in ef if ch.get(qn("r:id"))]
                efl.remove(ef)
                for rid in rids:
                    prs.part.drop_rel(rid)
        if len(efl) == 0:
            pres_el.remove(efl)
    return prs


def clear_slide(slide):
    """Remove every non-placeholder shape, dropping relationships of removed charts/pictures."""
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            continue
        rids = [v for el in shape._element.iter() for k, v in el.attrib.items()
                if k in (qn("r:id"), qn("r:embed"), qn("r:link"))]
        remove_shape(shape)
        for rid in rids:
            slide.part.drop_rel(rid)


# ---- charts -----------------------------------------------------------------
def _chart_no_frame(chart):
    cs = chart._chartSpace
    for old in cs.findall(qn("c:spPr")):
        cs.remove(old)
    spPr = etree.fromstring(
        '<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>")
    chart_el = cs.find(qn("c:chart"))
    chart_el.addnext(spPr)
    rc = cs.find(qn("c:roundedCorners"))
    if rc is None:
        rc = etree.Element(qn("c:roundedCorners"))
        cs.insert(0, rc)
    rc.set("val", "0")
    plot_area = chart_el.find(qn("c:plotArea"))
    for old in plot_area.findall(qn("c:spPr")):
        plot_area.remove(old)
    plot_area.append(copy.deepcopy(spPr))


def base_style(chart, y_min=None, y_max=None):
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = FONT_BODY
    chart.font.size = Pt(9)
    chart.font.color.rgb = rgb(AXIS)
    va = chart.value_axis
    va.has_major_gridlines = False
    va.has_minor_gridlines = False
    va.visible = False
    if y_min is not None:
        va.minimum_scale = y_min
    if y_max is not None:
        va.maximum_scale = y_max
    va.crosses = XL_AXIS_CROSSES.MINIMUM  # baseline sits at the bottom even when y_min < 0
    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = rgb(BASELINE)
    ca.format.line.width = Pt(0.75)
    ca.tick_labels.font.size = Pt(9)
    ca.tick_labels.font.name = FONT_BODY
    ca.tick_labels.font.color.rgb = rgb(AXIS)
    _chart_no_frame(chart)


def set_plot_area(chart, x, y, w, h):
    """Pin the inner plot area to fractions of the chart frame, so shapes placed next to the chart
    (annotation rows, CAGR arrows) can be aligned with the categories:
    category i centre = chart_x + chart_w * (x + w * (i + 0.5) / n)."""
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    old = plot_area.find(qn("c:layout"))
    if old is not None:
        plot_area.remove(old)
    layout = etree.fromstring(
        '<c:layout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:manualLayout>'
        '<c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
        f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>')
    plot_area.insert(0, layout)


def _label_point(series, idx, pos, color=BODY, size=9, bold=False, fill=None):
    """Value label on one point; fill = hex gives a solid label box (reference-chart style)."""
    dl = series.points[idx].data_label
    dl.position = pos
    dl.font.size = Pt(size)
    dl.font.name = FONT_BODY
    dl.font.bold = bold
    dl.font.color.rgb = rgb(color)
    dLbl = series.points[idx].data_label._dLbl
    if fill:
        spPr = dLbl.get_or_add_spPr()
        for ch in list(spPr):
            spPr.remove(ch)
        etree.SubElement(etree.SubElement(spPr, qn("a:solidFill")), qn("a:srgbClr")).set("val", fill)
        etree.SubElement(etree.SubElement(spPr, qn("a:ln")), qn("a:noFill"))
    for tag in ("c:showLegendKey", "c:showCatName", "c:showSerName", "c:showPercent", "c:showBubbleSize"):
        el = dLbl.find(qn(tag))
        if el is not None:
            el.set("val", "0")
    sv = dLbl.find(qn("c:showVal"))
    if sv is not None:
        sv.set("val", "1")


def _custom_label(series, idx, text, pos, color=BODY, size=9, bold=False):
    """Data label with free text (e.g. a stack total) instead of the point value."""
    dl = series.points[idx].data_label
    dl.has_text_frame = True
    dl.text_frame.text = text
    dl.position = pos
    for run in dl.text_frame.paragraphs[0].runs:
        style_run(run, size=size, color=color, bold=bold)


def add_line_chart(slide, x, y, w, h, categories, series, number_format="0.0", y_min=None, y_max=None,
                   name=None, baseline=None):
    """series: [{"name", "values", "color", "labels": {idx: XL_LABEL_POSITION},
                 optional "label_all": XL_LABEL_POSITION (every point), "label_fill": bool (solid box in
                 the series colour, white bold text), "dash": bool, "marker": bool (default True), "width": pt}]
    baseline: value where the category axis line crosses (e.g. 100 on an index chart); tick labels stay low."""
    cd = CategoryChartData(number_format=number_format)
    cd.categories = categories
    for s in series:
        cd.add_series(s["name"], s["values"])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Pt(x), Pt(y), Pt(w), Pt(h), cd)
    if name:
        gf.name = name
    chart = gf.chart
    base_style(chart, y_min, y_max)
    if baseline is not None:
        chart.value_axis.crosses_at = baseline
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    for ser, spec in zip(chart.plots[0].series, series):
        ser.smooth = False
        ser.format.line.color.rgb = rgb(spec["color"])
        ser.format.line.width = Pt(spec.get("width", 2.25))
        if spec.get("dash"):
            ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if spec.get("marker", True):
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 5
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(spec["color"])
            ser.marker.format.line.color.rgb = rgb(spec["color"])
        else:
            ser.marker.style = XL_MARKER_STYLE.NONE
        labels = spec.get("labels", {})
        if spec.get("label_all") is not None:
            labels = {i: spec["label_all"] for i in range(len(categories))}
        for idx, pos in labels.items():
            if spec.get("label_fill"):
                _label_point(ser, idx, pos, color="FFFFFF", size=spec.get("label_size", 9), bold=True,
                             fill=spec["color"])
            else:
                _label_point(ser, idx, pos)
    return gf


def add_column_chart(slide, x, y, w, h, categories, series, stacked=False, number_format="0.0",
                     gap_width=60, overlap=None, y_min=None, y_max=None, name=None, totals=None, total_pad=None):
    """series: [{"name", "values", "color", "label_color", "label_pos", "label_all": bool, "label_size",
                 optional "point_colors": [hex per category] (one series, bars in different colours),
                 "label_points": [idx, ...] (label only these points)}].
    totals (stacked only): list of label strings shown above each stack via a transparent top series."""
    cd = CategoryChartData(number_format=number_format)
    cd.categories = categories
    for s in series:
        cd.add_series(s["name"], s["values"])
    if totals:
        stack_max = max(sum(vals) for vals in zip(*[s["values"] for s in series]))
        cd.add_series("Total", [total_pad or stack_max * 0.08] * len(categories))
    kind = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Pt(x), Pt(y), Pt(w), Pt(h), cd)
    if name:
        gf.name = name
    chart = gf.chart
    base_style(chart, y_min, y_max)
    plot = chart.plots[0]
    plot.gap_width = gap_width
    plot.overlap = overlap if overlap is not None else (100 if stacked else -10)
    for ser, spec in zip(plot.series, series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(spec["color"])
        ser.format.line.color.rgb = rgb("FFFFFF")
        ser.format.line.width = Pt(0.75)
        ser.invert_if_negative = False
        for i, col in enumerate(spec.get("point_colors") or []):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(col)
            pt.format.line.color.rgb = rgb("FFFFFF")
        points = spec.get("label_points")
        if points is None and spec.get("label_all", True):
            points = range(len(categories))
        if points:
            pos = spec.get("label_pos", XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END)
            for idx in points:
                _label_point(ser, idx, pos, color=spec.get("label_color", BODY), size=spec.get("label_size", 9))
    if totals:
        tser = plot.series[len(series)]
        tser.format.fill.background()
        tser.format.line.fill.background()
        for idx, text in enumerate(totals):
            _custom_label(tser, idx, text, XL_LABEL_POSITION.INSIDE_BASE, bold=True)
    return gf
