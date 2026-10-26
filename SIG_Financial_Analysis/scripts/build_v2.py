"""Build v2 of the SIG financial analysis deck: completes slides 5, 6 and 13.

Reads v1 (never modified) and the WP workbook; writes v2. Run:
    python build_v2.py
"""
import shutil
import sys
from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.enum.chart import XL_LABEL_POSITION as POS
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

sys.path.insert(0, str(Path(__file__).parent))
from skha_helpers import (AXIS, BODY, FONT_BODY, GREY, IDLE, INK, NAVY, PANEL_GREY, RED, SKHA_BLUE, SKY, TINT,  # noqa: E402
                          WATCH, add_banner, add_column_chart, add_column_header, add_kpi_table, add_legend,
                          add_line_chart, add_rect, add_takeaway_box, add_text, fill_text_frame, remove_shape,
                          replace_runs_text, rgb)

SRC = Path(r"D:\Download\Financial Analysis SIG (SMGR) v1.pptx")
OUT = Path(r"D:\Download\Financial Analysis SIG (SMGR) v2.pptx")
XL = Path(r"D:\Download\WP - Financial Analysis SIG (SMGR) v0.xlsx")
YEARS = ["FY21", "FY22", "FY23", "FY24", "FY25"]

# ---- data -------------------------------------------------------------------
wb = openpyxl.load_workbook(XL, data_only=True)


def ratio(no):
    """Ratios v2 row by 'No': returns (SIG FY21-25, INTP FY21-25)."""
    ws = wb["Ratios v2"]
    for r in range(5, 80):
        v = ws.cell(r, 1).value
        if isinstance(v, (int, float)) and int(v) == no:
            sig = [ws.cell(r, c).value for c in range(5, 10)]
            intp = [ws.cell(r, c).value for c in range(12, 17)]
            return sig, intp
    raise KeyError(no)


def row_values(sheet, label, n=5, start_offset=1, startswith=True):
    """First row whose label cell matches; returns the next n numeric cells to its right."""
    ws = wb[sheet]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and (c.value.strip().startswith(label) if startswith else label in c.value):
                vals = []
                col = c.column + start_offset
                while len(vals) < n and col < c.column + 20:
                    v = ws.cell(c.row, col).value
                    if isinstance(v, (int, float)):
                        vals.append(v)
                    col += 1
                if len(vals) == n:
                    return vals
    raise KeyError(f"{sheet}: {label}")


def header_block(sheet, header_row_label, header):
    """Values under a header cell for the 5 rows that follow it."""
    ws = wb[sheet]
    for row in ws.iter_rows():
        for c in row:
            if c.value == header:
                return [ws.cell(c.row + 1 + i, c.column).value for i in range(5)]
    raise KeyError(header)


D = {}
D["ebitda_m"] = ratio(3)
D["net_m"] = ratio(5)
D["roce"] = ratio(9)
D["de"] = ratio(27)
D["ibd_ebitda"] = ratio(31)
D["tie"] = ratio(33)
D["current"] = ratio(36)
D["quick"] = ratio(37)
D["ocf"] = ratio(40)
D["cash_due"] = ratio(41)
D["at"] = ratio(11)
D["fat"] = ratio(16)
D["dpo"] = ratio(47)
D["ccc"] = ratio(48)
D["ebitda_tn"] = [v / 1e6 for v in row_values("IS (Build)", "EBITDA (EBIT + D&A)")]
D["ibd_tn"] = [v / 1e6 for v in row_values("BS (Presentation)", "Memo: Utang berbunga", startswith=False)]
D["rev_ton"] = row_values("IS (Build)", "Pendapatan / ton (Rp ribu)")
D["vol_total"] = [v / 1e3 for v in row_values("IS (Build)", "Total volume semen & terak")]
D["vol_regional"] = [v / 1e3 for v in row_values("IS (Build)", "Volume penjualan regional/ekspor SIG")]
D["vol_dom_cagr"] = row_values("IS (Build)", "Volume penjualan domestik SIG", n=6)[5]
D["mkt_dom_cagr"] = row_values("IS (Build)", "Permintaan semen domestik nasional", n=6)[5]
D["clinker_rev"] = [v / 1e6 for v in row_values("IS (Build)", "Terak (klinker)")]
D["energy_ton_cagr"] = row_values("IS (Build)", "  Bahan bakar & energi / ton", n=6, startswith=False)[5]
D["bulk_share"] = row_values("IS (Build)", "  Pangsa pasar – curah", n=2, startswith=False)
D["share_sig"] = header_block("Business Context", None, "% SIG")
D["share_intp"] = header_block("Business Context", None, "%INTP")
D["share_oth"] = header_block("Business Context", None, "% Others")
D["hba"] = row_values("Business Context", "Average HBA (USD/ton)")
D["intp_rev_tn"] = row_values("Peer_INTP", "Pendapatan neto")[4] / 1e6
D["intp_vol"] = row_values("INTP_3FS", "Volume penjualan semen & klinker")[4] / 1e3
D["sig_rev_tn"] = row_values("IS (Build)", "Pendapatan")[4] / 1e6
D["sig_capacity"] = header_block("WP", None, "Capacity")[4]
INTP_CAPACITY = 33.5  # Mt, INTP FY2025 press release (14 owned + 2 leased plants + 1 grinding mill)

# industry capacity vs absorption (Business Context, "Chart data – Capacity vs Absorption")
bc = wb["Business Context"]
hdr = {bc.cell(61, c).value: c for c in range(15, 23)}
col_dd = [c for v, c in hdr.items() if v == "Domestic Demand"][0]
IND = {k: [bc.cell(62 + i, col_dd + off).value for i in range(5)]
       for k, off in [("demand", 0), ("export", 1), ("idle", 2), ("capacity", 3), ("util", 4)]}

for k, v in {**D, **IND}.items():
    vals = v if isinstance(v, (list, tuple)) else [v]
    flat = [x for item in vals for x in (item if isinstance(item, list) else [item])]
    assert all(isinstance(x, (int, float)) for x in flat), (k, v)


def f1(x):
    return f"{x:.1f}"


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


def xr(x, d=2):
    return f"{x:.{d}f}x"


# ---- deck -------------------------------------------------------------------
shutil.copyfile(SRC, OUT)
prs = Presentation(OUT)


def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(idx)


def set_title(slide, text):
    replace_runs_text(ph(slide, 0), text)


def set_source(slide, text):
    fill_text_frame(ph(slide, 10).text_frame, [text], size=8, color="7F7F7F")


def remove_yellow(slide):
    for s in list(slide.shapes):
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.fill.type == 1:
            fc = s.fill.fore_color
            if fc.type == MSO_COLOR_TYPE.RGB and str(fc.rgb) == "FFC000":
                remove_shape(s)


def find(shapes, name):
    for s in shapes:
        if s.name == name:
            return s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            hit = find(s.shapes, name)
            if hit is not None:
                return hit
    return None


# ============================ SLIDE 5 ==========================================
s5 = prs.slides[4]
nm_sig, nm_intp = D["net_m"]
set_title(s5, "Revenue is flat for both SIG and Indocement, but only SIG\u2019s profitability collapsed as costs "
              f"outgrew revenue and net margin fell from {pct(nm_sig[0])} to {pct(nm_sig[4])}")

# left chart: recolour cost stack to template ramp, quiet labels
ch = find(s5.shapes, "Revenue vs cost SIG").chart
for ser, color in zip(ch.plots[0].series, [TINT, SKY, NAVY]):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(color)
ch.legend.font.color.rgb = rgb(AXIS)
ch.legend.font.size = 9 * 12700
ch.category_axis.tick_labels.font.color.rgb = rgb(AXIS)

# KPI panel: SIG navy, INTP red (comparator)
kpi = find(s5.shapes, "Rectangle 27")
for p in kpi.text_frame.paragraphs:
    for r in p.runs:
        c = r.font.color
        if c.type == MSO_COLOR_TYPE.RGB and str(c.rgb) in ("C00000", "1E64B0"):
            r.font.color.rgb = rgb(NAVY if str(c.rgb) == "C00000" else RED)

# right column header
em_sig, em_intp = D["ebitda_m"]
ro_sig, ro_intp = D["roce"]
gap_em = (em_sig[4] - em_intp[4]) * 100
replace_runs_text(find(s5.shapes, "Rectangle 31"),
                  f"SIG\u2019s margins matched INTP\u2019s in FY21; by FY25 its EBITDA margin trails by "
                  f"{abs(gap_em):.1f} pp and ROCE is a quarter of INTP\u2019s")
replace_runs_text(find(s5.shapes, "Rectangle 30"), "EBITDA margin and ROCE, %, FY21\u2013FY25")
assert 0.2 <= ro_sig[4] / ro_intp[4] <= 0.3  # "about a quarter"

RX, RW = 539.2, 377.1
add_legend(s5, RX, 177, [("SIG", NAVY, "line"), ("INTP", RED, "line")])
for (top, label, gap_text, sig, intp, ymin, ymax, sig_pos, intp_pos) in [
    (191, "EBITDA margin, %", f"FY25 gap: \u2212{abs(gap_em):.1f} pp", em_sig, em_intp, 0.08, 0.28,
     {0: POS.ABOVE, 4: POS.BELOW}, {0: POS.BELOW, 4: POS.ABOVE}),
    (319, "ROCE, %", f"FY25: {pct(ro_sig[4])} vs {pct(ro_intp[4])}", ro_sig, ro_intp, -0.035, 0.13,
     {0: POS.BELOW, 4: POS.BELOW}, {0: POS.ABOVE, 4: POS.ABOVE}),
]:
    add_text(s5, RX, top, 200, 12, label, size=9.75)
    add_text(s5, RX + RW - 180, top, 180, 12, [[(gap_text, {"bold": True})]], size=9.75, align="r")
    add_line_chart(s5, RX, top + 10, RW, 116, YEARS,
                   [{"name": "SIG", "values": sig, "color": NAVY, "labels": sig_pos},
                    {"name": "INTP", "values": intp, "color": RED, "labels": intp_pos}],
                   number_format="0.0%", y_min=ymin, y_max=ymax, name=f"Chart {label}")

# bottom box -> takeaway
box = find(s5.shapes, "Rectangle 6")
tf = box.text_frame
tf.margin_left = tf.margin_right = 12 * 12700
fill_text_frame(tf, [
    [("Same market, different outcome: ", {"bold": True}),
     "INTP kept its margins close to FY21 levels, so SIG\u2019s decline is company-specific, not market-wide."],
    [("Key question: ", {"bold": True}), "what is driving SIG\u2019s profitability down?"],
], size=12)
set_source(s5, "Note: SIG net margin uses profit attributable to owners; ROCE = EBIT \u00f7 (total assets \u2212 current "
               "liabilities). Source: SIG and INTP audited consolidated financial statements FY2021\u20132025; Skha analysis")

# ============================ SLIDE 6 ==========================================
s6 = prs.slides[5]
remove_yellow(s6)
set_title(s6, "Leverage, liquidity and efficiency are not the root cause: their weak spots trace back to falling "
              "EBITDA and idle capacity, so the deep dive focuses on profitability")
COLS = [40.6, 338.0, 635.3]
CW = 283.5
KPI_W = [118.5, 55, 55, 55]
KPI_HDR = ["SIG FY21", "SIG FY25", "INTP FY25"]
DOWN = "\u25bc"


def three(key, fmt):
    s, i = D[key]
    return [fmt(s[0]), fmt(s[4]), fmt(i[4])]


ibd_e = D["ibd_ebitda"][0]
pillars = [
    {
        "x": COLS[0], "name": "Leverage", "verdict": "Second-order risk",
        "headline": f"Debt halved, but so did EBITDA, so debt/EBITDA is still ~{ibd_e[4]:.1f}x",
        "caption": "Interest-bearing debt and EBITDA, Rp tn",
        "legend": [("Interest-bearing debt", NAVY, "square"), ("EBITDA", SKY, "square")],
        "rows": [
            ("Debt-to-equity", three("de", xr)),
            ("Interest-bearing debt \u00f7 EBITDA", three("ibd_ebitda", xr)),
            ("EBIT \u00f7 interest", [xr(D["tie"][0][0], 1), (xr(D["tie"][0][4], 1), DOWN), xr(D["tie"][1][4], 1)]),
        ],
    },
    {
        "x": COLS[1], "name": "Liquidity", "verdict": "Watch item",
        "headline": f"Current ratio holds at ~{D['current'][0][4]:.1f}x, but cash buffers are far thinner "
                    "than INTP\u2019s",
        "caption": "Current ratio, x",
        "legend": [("SIG", NAVY, "line"), ("INTP", RED, "line")],
        "rows": [
            ("Cash \u00f7 debt due \u22641 year", three("cash_due", xr)),
            ("Operating cash flow \u00f7 current liabilities",
             [xr(D["ocf"][0][0]), (xr(D["ocf"][0][4]), DOWN), xr(D["ocf"][1][4])]),
            ("Quick ratio", three("quick", xr)),
        ],
    },
    {
        "x": COLS[2], "name": "Efficiency", "verdict": "Capped by idle capacity",
        "headline": "Asset turnover is flat; the gap to INTP reflects a heavy, under-used plant base",
        "caption": "Fixed asset turnover, x",
        "legend": [("SIG", NAVY, "line"), ("INTP", RED, "line")],
        "rows": [
            ("Asset turnover", three("at", xr)),
            ("Cash conversion cycle, days", three("ccc", lambda v: f"{v:.0f}")),
            ("Payables days", three("dpo", lambda v: f"{v:.0f}")),
        ],
    },
]
for pl in pillars:
    x = pl["x"]
    add_banner(s6, x, 114.1, CW, 22, pl["name"], pl["verdict"])
    add_column_header(s6, x, 141, CW, pl["headline"], pl["caption"])
    add_legend(s6, x, 193, pl["legend"])
    kpi_y = 323
    add_kpi_table(s6, x, kpi_y, KPI_W, KPI_HDR, pl["rows"], row_h=22)

add_column_chart(s6, COLS[0], 205, CW, 114, YEARS,
                 [{"name": "Interest-bearing debt", "values": D["ibd_tn"], "color": NAVY},
                  {"name": "EBITDA", "values": D["ebitda_tn"], "color": SKY}],
                 gap_width=70, y_min=0, y_max=max(D["ibd_tn"]) * 1.18, name="Chart leverage")
cs, ci = D["current"]
add_line_chart(s6, COLS[1], 205, CW, 114, YEARS,
               [{"name": "SIG", "values": cs, "color": NAVY, "labels": {0: POS.BELOW, 4: POS.BELOW}},
                {"name": "INTP", "values": ci, "color": RED, "labels": {0: POS.ABOVE, 4: POS.ABOVE}}],
               number_format="0.00", y_min=0.3, y_max=2.8, name="Chart liquidity")
fs, fi = D["fat"]
add_line_chart(s6, COLS[2], 205, CW, 114, YEARS,
               [{"name": "SIG", "values": fs, "color": NAVY, "labels": {0: POS.BELOW, 4: POS.BELOW}},
                {"name": "INTP", "values": fi, "color": RED, "labels": {0: POS.ABOVE, 4: POS.ABOVE}}],
               number_format="0.00", y_min=0.4, y_max=1.25, name="Chart efficiency")

add_takeaway_box(s6, 40.6, 424, 878.2, 62, [
    [("None of the three is the root cause. ", {"bold": True}),
     "Leverage and liquidity weaken only because EBITDA falls, and efficiency is capped by idle capacity, "
     "so the next section breaks down profitability."],
])
set_source(s6, "Note: interest-bearing debt includes temporary syirkah funds and leases; SIG equity = equity "
               "attributable to owners. Source: SIG and INTP audited financial statements FY2021\u20132025; Skha analysis")

# ============================ SLIDE 13 =========================================
s13 = prs.slides[12]
remove_yellow(s13)
set_title(s13, "Flat demand against ~120 Mt of capacity, a coal-price shock and rivals gaining share squeezed the "
               "industry; SIG defended price and lost share")
LX, LW = 40.6, 498.7
RXC, RWC = 550.9, 368.5

# --- panel A: market
add_column_header(s13, LX, 114.1, LW,
                  f"Demand has been flat at ~{IND['demand'][4]:.0f} Mt against ~{IND['capacity'][4]:.0f} Mt of "
                  "capacity, while energy costs spiked in 2022\u201323",
                  "Installed capacity vs. domestic demand and export, Mt")
add_legend(s13, LX, 166, [("Domestic demand", SKHA_BLUE, "square"), ("Export", SKY, "square"),
                          ("Idle capacity", IDLE, "square")], gap=10)
add_column_chart(s13, LX, 177, 248, 121, YEARS,
                 [{"name": "Domestic demand", "values": IND["demand"], "color": SKHA_BLUE, "label_color": "FFFFFF"},
                  {"name": "Export", "values": IND["export"], "color": SKY, "label_all": False},
                  {"name": "Idle capacity", "values": IND["idle"], "color": IDLE, "label_color": BODY}],
                 stacked=True, gap_width=45, y_min=0, y_max=125, name="Chart industry capacity")
util_lo, util_hi = min(IND["util"]) * 100, max(IND["util"]) * 100
add_text(s13, LX + 258, 168, LW - 258, 130, [
    {"runs": [("Oversupply: ", {"bold": True}),
              f"total utilization stuck at {util_lo:.0f}\u2013{util_hi:.0f}% (domestic + export); "
              f"~{IND['idle'][4]:.0f} Mt of capacity idle every year."], "bullet": True, "space_after": 4},
    {"runs": [("Regions: ", {"bold": True}),
              "Java takes ~51% of national demand; Kalimantan fell 16% in 2025 as IKN construction slowed."],
     "bullet": True, "space_after": 4},
    {"runs": [("Costs: ", {"bold": True}),
              f"coal reference price (HBA) jumped from USD {D['hba'][0]:.0f} to {D['hba'][1]:.0f}/t in 2022; "
              f"SIG energy cost per ton rose {D['energy_ton_cagr'] * 100:.1f}% a year. Carbon tax is still not "
              "levied, a future risk."], "bullet": True},
], size=9.75)

# --- panel B: competitors
sh_s, sh_i, sh_o = D["share_sig"], D["share_intp"], D["share_oth"]
gain_intp = (sh_i[4] - sh_i[0]) * 100
em_ratio = em_intp[4] / em_sig[4]
add_column_header(s13, LX, 306, LW,
                  f"INTP gained {gain_intp:.1f} pp of share since FY21 and earns an EBITDA margin "
                  f"~{em_ratio:.1f}x SIG\u2019s",
                  "Key metrics, FY2025 unless stated")
TB_W = [190, 102.9, 102.9, 102.9]
ty = 357
cx = LX + TB_W[0]
for label, color, w in [("SIG", NAVY, TB_W[1]), ("INTP", RED, TB_W[2]), ("Others", GREY, TB_W[3])]:
    tw = len(label) * 5.2
    add_rect(s13, cx + w - tw - 9, ty + 4.5, 6, 6, fill=color)
    add_text(s13, cx, ty, w, 14, [[(label, {"bold": True})]], size=9.75, align="r", anchor="m")
    cx += w
cmp_rows = [
    ("Domestic market share, FY21 \u2192 FY25",
     [f"{pct(sh_s[0])} \u2192 {pct(sh_s[4])}", f"{pct(sh_i[0])} \u2192 {pct(sh_i[4])}",
      f"{pct(sh_o[0])} \u2192 {pct(sh_o[4])}"]),
    ("Installed capacity, Mt", [f1(D["sig_capacity"]), f1(INTP_CAPACITY), "\u2013"]),
    ("Sales volume, Mt", [f1(D["vol_total"][4]), f1(D["intp_vol"]), "\u2013"]),
    ("Revenue, Rp tn", [f1(D["sig_rev_tn"]), f1(D["intp_rev_tn"]), "\u2013"]),
    ("EBITDA margin", [pct(em_sig[4]), pct(em_intp[4]), "\u2013"]),
]
add_kpi_table(s13, LX, ty, TB_W, ["", "", ""], cmp_rows, row_h=21, header_h=14, bold_col=None)

# --- panel C: SIG response
add_column_header(s13, RXC, 114.1, RWC,
                  "SIG defended price over volume and lost share; its 2025 turnaround is too early to judge",
                  "SIG\u2019s strategic moves since FY21 and their results")
reg_share = D["vol_regional"][4] / D["vol_total"][4]
clinker_yoy = D["clinker_rev"][4] / D["clinker_rev"][3] - 1
moves = [
    ("Held price over volume. ",
     f"Revenue per ton rose from Rp{D['rev_ton'][0]:.0f}k to Rp{D['rev_ton'][4]:.0f}k "
     f"(+{((D['rev_ton'][4] / D['rev_ton'][0]) ** 0.25 - 1) * 100:.1f}% a year) while volume fell from "
     f"{D['vol_total'][0]:.1f} to {D['vol_total'][4]:.1f} Mt."),
    ("Filled idle kilns with regional and export sales. ",
     f"~{D['vol_regional'][4]:.1f} Mt in FY25 ({reg_share * 100:.0f}% of volume); clinker revenue rose "
     f"{clinker_yoy * 100:.0f}% in FY25 to Rp{D['clinker_rev'][4]:.1f} tn."),
    ("Ran a multi-brand portfolio with fighting brands. ",
     "7+ brands across micro-markets; fighting-brand prices were raised 2.2% in Jun\u2013Sep 2025."),
]
my = 168
for i, (lead, body) in enumerate(moves, 1):
    circ = add_rect(s13, RXC, my + 1, 15, 15, fill=SKHA_BLUE, shape=MSO_SHAPE.OVAL)
    fill_text_frame(circ.text_frame, [[(str(i), {"bold": True})]], size=9, color="FFFFFF", align="c")
    add_text(s13, RXC + 22, my, RWC - 22, 42, [[(lead, {"bold": True}), body]], size=9.75)
    my += 45
add_rect(s13, RXC, 304, RWC, 52, fill=PANEL_GREY)
add_text(s13, RXC + 8, 304, RWC - 16, 52, [[
    ("Result: ", {"bold": True}),
    f"domestic share fell from {pct(sh_s[0])} to {pct(sh_s[4])}, steepest in bulk "
    f"({pct(D['bulk_share'][0])} \u2192 {pct(D['bulk_share'][1])}, FY24\u201325); SIG domestic volume fell "
    f"{abs(D['vol_dom_cagr']) * 100:.1f}% a year vs. a market down {abs(D['mkt_dom_cagr']) * 100:.1f}%."]],
    size=9.75, anchor="m")
add_text(s13, RXC, 364, 230, 16, [[("2025 turnaround (since Jul 2025)", {"bold": True})]], size=10.5,
         color=INK, anchor="m")
chip = add_rect(s13, RXC + RWC - 112, 365, 112, 14, fill=WATCH, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
fill_text_frame(chip.text_frame, [[("Too early to evaluate", {"bold": True})]], size=8.5, align="c")
add_text(s13, RXC, 384, RWC, 110, [
    {"runs": [("Retail business model: ", {"bold": True}),
              "sales and outbound logistics handed to the cement subsidiary in each area"], "bullet": True},
    {"runs": [("Derivatives & portfolio: ", {"bold": True}), "mortar, interlocking bricks, ready-mix"],
     "bullet": True},
    {"runs": [("Cost leadership: ", {"bold": True}),
              "outbound supply chain, inter-plant logistics, fixed and variable costs"], "bullet": True},
    {"runs": [("Streamlining: ", {"bold": True}), "40 entities to 12, in execution from Q2 2026"], "bullet": True},
    {"runs": [("Early read: ", {"bold": True}), "market share 46.5% \u2192 46.8% (1H25 \u2192 9M25)"],
     "bullet": True},
], size=9.75, space_after=1.5)
set_source(s13, "Source: SIG annual reports, financial statements and 9M25 earnings call (Nov 2025); INTP FY2025 press "
                "release and 2026 public expose (Kemenperin, BPS); ESDM (HBA); BCA Sekuritas (May 2026); DDTC (Oct "
                "2026); Skha analysis")

prs.save(OUT)
print("saved", OUT)
for k in ("ebitda_m", "roce", "current", "fat"):
    print(k, [round(v, 3) for v in D[k][0]], [round(v, 3) for v in D[k][1]])
print("ibd_tn", [round(v, 1) for v in D["ibd_tn"]], "ebitda_tn", [round(v, 1) for v in D["ebitda_tn"]])
print("industry", {k: [round(v, 3) for v in vals] for k, vals in IND.items()})
print("shares", sh_s, sh_i, sh_o, "bulk", D["bulk_share"], "rev/ton", D["rev_ton"])
